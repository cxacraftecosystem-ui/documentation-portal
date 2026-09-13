"""The edit-after-answers rule, as tests, with no database anywhere near it.

WHAT A FAKE CAN AND CANNOT PROVE HERE, stated because the boundary is the whole design of this file.

It proves WHICH WRITES THIS MODULE ISSUES AND IN WHAT ORDER — that a reworded answered question is
written as create-then-retire and not as an update, that every section is negated before anything is
renumbered, that a question with an answer is never handed to ``delete``, and that NO SECTION IS EVER
HANDED TO ``delete`` AT ALL.

It does NOT prove that Postgres accepts them, and it cannot prove the consequence of the one call it
forbids. The unique on ``(questionnaireId, sortOrder)``, the ON DELETE RESTRICT on
``QuestionnaireResponse.questionId`` and ``QuestionnaireQuestion.sectionId``, and the ON DELETE
CASCADE on ``QuestionnaireSectionStatus.sectionId`` are real constraints this fake does not enforce.
A section delete would pass every assertion here and destroy every admin completion verdict for that
section in production. The only way to exercise those is ``prisma migrate deploy`` against a local
Postgres and the by-hand walkthrough in the pull request.

So the assertions below are about CALLS, deliberately and exclusively. Where a rule's consequence is
invisible from here, the assertion says which table it is invisible about, so the next reader knows
what this file is not telling them.
"""

import asyncio
from types import SimpleNamespace
from typing import Any

import pytest

import app.api.routes.questionnaire as workbook
from app.services.questionnaire_xlsx import (
    ParsedQuestion,
    ParsedQuestionnaire,
    ParsedSection,
)

# =================================================================================================
# The fake
# =================================================================================================


class _Delegate:
    """One Prisma model delegate, recording every call made against it.

    It implements the QUERIES, never the DECISIONS: ``find_many`` filters on the handful of clauses
    this module actually issues, so a test of "which questions did it read" is a test of the module's
    where-clause rather than of a helper that was told the answer.
    """

    def __init__(
        self,
        name: str,
        journal: list[tuple[str, str, dict]],
        rows: list[Any],
        defaults: dict[str, Any] | None = None,
    ) -> None:
        self.name = name
        self.journal = journal
        self.rows = rows
        # THE SCHEMA'S COLUMN DEFAULTS, APPLIED ON INSERT THE WAY POSTGRES APPLIES THEM. Without
        # these the row a `create` hands back is missing every column the caller did not set, and the
        # module reads `questionnaire.version` off it immediately — so a fake without defaults fails
        # with an AttributeError that looks like a bug in the module rather than in the fake.
        self.defaults = defaults or {}
        self._next = 0

    def _record(self, op: str, **kwargs: Any) -> None:
        self.journal.append((self.name, op, kwargs))

    def _mint(self, prefix: str) -> str:
        self._next += 1
        return f"{prefix}-new-{self._next}"

    async def create(self, data: dict[str, Any]) -> Any:
        self._record("create", data=data)
        fields = {**self.defaults, **data}
        fields.setdefault("id", self._mint(self.name))
        row = SimpleNamespace(**fields)
        self.rows.append(row)
        return row

    async def update(self, where: dict[str, Any], data: dict[str, Any]) -> Any:
        self._record("update", where=where, data=data)
        for row in self.rows:
            if row.id == where.get("id"):
                for key, value in data.items():
                    setattr(row, key, value)
                return row
        return SimpleNamespace(id=where.get("id"), **data)

    async def update_many(self, where: dict[str, Any], data: dict[str, Any]) -> int:
        self._record("update_many", where=where, data=data)
        return 0

    async def delete(self, where: dict[str, Any]) -> Any:
        self._record("delete", where=where)
        self.rows = [row for row in self.rows if row.id != where.get("id")]
        return None

    async def count(self, where: dict[str, Any] | None = None) -> int:
        self._record("count", where=where)
        return self.counts if isinstance(getattr(self, "counts", None), int) else 0

    async def find_unique(self, where: dict[str, Any], **_: Any) -> Any:
        self._record("find_unique", where=where)
        return next((row for row in self.rows if row.id == where.get("id")), None)

    async def find_many(self, where: dict[str, Any] | None = None, **kwargs: Any) -> list[Any]:
        self._record("find_many", where=where, **kwargs)
        rows = list(self.rows)
        where = where or {}
        for key, clause in where.items():
            if isinstance(clause, dict) and "in" in clause:
                rows = [r for r in rows if getattr(r, key, None) in clause["in"]]
            else:
                rows = [r for r in rows if getattr(r, key, None) == clause]
        take = kwargs.get("take")
        if kwargs.get("order") and rows:
            order = kwargs["order"]
            first = order[0] if isinstance(order, list) else order
            field, direction = next(iter(first.items()))
            rows.sort(key=lambda r: getattr(r, field, 0), reverse=direction == "desc")
        return rows[:take] if take else rows


class _Db:
    """``db`` for the duration of one test. Any delegate not named here is a mistake, loudly."""

    def __init__(self, **delegates: _Delegate) -> None:
        self._delegates = delegates

    def __getattr__(self, name: str) -> Any:
        try:
            return self._delegates[name]
        except KeyError as exc:  # pragma: no cover - only read out of a failure
            raise AssertionError(
                f"the import reached for db.{name}, which this test did not model. If that is "
                "correct, model it; if it is not, the module is touching a table it should not."
            ) from exc


@pytest.fixture
def world(monkeypatch: pytest.MonkeyPatch):
    """A questionnaire with two sections and three questions, and a journal of every write.

    The module does ``from app.core.db import db``, so it holds its OWN reference and patching the
    source module alone would miss it. Patched by identity on the module object, exactly as
    ``test_permission_matrix.py`` does.
    """
    journal: list[tuple[str, str, dict]] = []

    questionnaire = SimpleNamespace(
        id="qnr-1",
        title="2nd Craft Toolkit Workshop",
        description="The instrument",
        version=4,
        sortOrder=1,
        isActive=True,
        isDefault=True,
        sourceFilename=None,
        createdAt="2026-01-01",
    )
    sections = [
        SimpleNamespace(
            id="sec-a", questionnaireId="qnr-1", code="A", title="About", sortOrder=1, isActive=True
        ),
        SimpleNamespace(
            id="sec-b",
            questionnaireId="qnr-1",
            code="B",
            title="Materials",
            sortOrder=2,
            isActive=True,
        ),
    ]
    questions = [
        SimpleNamespace(
            id="q-1",
            questionnaireId="qnr-1",
            sectionId="sec-a",
            sectionCode="A",
            sectionTitle="About",
            prompt="How long have you practised this craft?",
            helpText=None,
            isRequired=False,
            sortOrder=1,
            isActive=True,
            retiredAt=None,
            supersededById=None,
        ),
        SimpleNamespace(
            id="q-2",
            questionnaireId="qnr-1",
            sectionId="sec-a",
            sectionCode="A",
            sectionTitle="About",
            prompt="Who taught you?",
            helpText=None,
            isRequired=False,
            sortOrder=2,
            isActive=True,
            retiredAt=None,
            supersededById=None,
        ),
        SimpleNamespace(
            id="q-3",
            questionnaireId="qnr-1",
            sectionId="sec-b",
            sectionCode="B",
            sectionTitle="Materials",
            prompt="Where do you buy your yarn?",
            helpText=None,
            isRequired=False,
            sortOrder=1,
            isActive=True,
            retiredAt=None,
            supersededById=None,
        ),
    ]

    delegates = {
        "questionnaire": _Delegate(
            "questionnaire",
            journal,
            [questionnaire],
            # prisma/schema.prisma: version @default(1), isDefault/isActive @default, the rest null.
            defaults={
                "version": 1,
                "isActive": True,
                "isDefault": False,
                "description": None,
                "sourceFilename": None,
                "createdById": None,
                "createdAt": "2026-09-13",
            },
        ),
        "questionnairesection": _Delegate(
            "questionnairesection", journal, sections, defaults={"isActive": True}
        ),
        "questionnairequestion": _Delegate(
            "questionnairequestion",
            journal,
            questions,
            defaults={
                "isActive": True,
                "isRequired": False,
                "helpText": None,
                "retiredAt": None,
                "supersededById": None,
                "sectionId": None,
            },
        ),
        "questionnaireresponse": _Delegate("questionnaireresponse", journal, []),
        "questionnairesectionstatus": _Delegate("questionnairesectionstatus", journal, []),
        "assignedtask": _Delegate("assignedtask", journal, []),
    }
    db = _Db(**delegates)
    monkeypatch.setattr(workbook, "db", db)
    # `row` holds the Questionnaire ROW; `questionnaire` is its DELEGATE, out of `delegates` below.
    # Naming them apart matters here: a test that reached for `world.questionnaire.version` and got a
    # delegate would read None and assert nothing.
    return SimpleNamespace(
        db=db,
        journal=journal,
        row=questionnaire,
        sections=sections,
        questions=questions,
        **delegates,
    )


def _answered(world, question_id: str, *, text: str = "Forty years.", notes: str | None = None):
    """Record a response row against a question, as an interview would."""
    world.questionnaireresponse.rows.append(
        SimpleNamespace(
            id=f"resp-{question_id}", questionId=question_id, answerText=text, notes=notes
        )
    )


def _parsed(*sections, questionnaire_id: str | None = None, **kwargs) -> ParsedQuestionnaire:
    return ParsedQuestionnaire(
        sections=list(sections),
        sheet="Questionnaire",
        questionnaireId=questionnaire_id,
        **kwargs,
    )


def _section(code: str, title: str, *questions) -> ParsedSection:
    return ParsedSection(code=code, title=title, questions=list(questions))


def _question(prompt: str, row: int, **kwargs) -> ParsedQuestion:
    return ParsedQuestion(prompt=prompt, row=row, **kwargs)


def _calls(world, name: str, op: str) -> list[dict]:
    return [kwargs for model, operation, kwargs in world.journal if model == name and operation == op]


def _run(coro):
    return asyncio.run(coro)


# =================================================================================================
# Creating a questionnaire from a workbook
# =================================================================================================


def test_a_new_workbook_writes_sections_then_questions_in_sheet_order(world):
    parsed = _parsed(
        _section("A", "About", _question("How long?", 2), _question("Who taught you?", 3)),
        _section("B", "Materials", _question("Where do you buy yarn?", 4)),
    )
    _run(workbook.create_from_parsed(parsed, created_by_id="admin-1"))

    order = [
        (model, kwargs["data"].get("code") or kwargs["data"].get("prompt"))
        for model, op, kwargs in world.journal
        if op == "create" and model in ("questionnairesection", "questionnairequestion")
    ]
    assert order == [
        ("questionnairesection", "A"),
        ("questionnairequestion", "How long?"),
        ("questionnairequestion", "Who taught you?"),
        ("questionnairesection", "B"),
        ("questionnairequestion", "Where do you buy yarn?"),
    ]


def test_helpText_and_isRequired_reach_the_create_call(world):
    """THE MIGRATION'S WHOLE JUSTIFICATION. The pro-forma writes a "Help text" column and a
    "Required" column and the parser reads both back; if they do not reach the create, the app asked
    an admin a question and then threw the answer away, with no row-level failure to report."""
    parsed = _parsed(
        _section(
            "A",
            "About",
            _question("How long?", 2, helpText="Count from the first unsupervised year.", isRequired=True),
        )
    )
    _run(workbook.create_from_parsed(parsed))

    data = _calls(world, "questionnairequestion", "create")[0]["data"]
    assert data["helpText"] == "Count from the first unsupervised year."
    assert data["isRequired"] is True


def test_a_help_text_over_the_cap_is_clipped_and_reported(world):
    """A CLIP WITH NO PROBLEM ROW IS THE SILENT DROP. The editor's own schema caps help text, so an
    over-long import would 422 every later PATCH from the builder with a message about a limit the
    admin never met and cannot see. Clipping makes the two agree; reporting is what stops the clip
    being invisible."""
    parsed = _parsed(
        _section("A", "About", _question("How long?", 7, helpText="x" * 5000))
    )
    _, report = _run(workbook.create_from_parsed(parsed))

    data = _calls(world, "questionnairequestion", "create")[0]["data"]
    assert len(data["helpText"]) == workbook.MAX_HELP_CHARS
    problem = next(p for p in report["problems"] if "help text" in p["reason"])
    assert problem["severity"] == "warning"
    assert problem["row"] == 7


def test_every_question_create_carries_sectionCode_and_sectionTitle(world):
    """THE FIELD-ONLY DENORMALISED NOT NULL COLUMNS, plus ``questionnaireId``, which is NOT NULL too.
    Omitting any of them is a NOT NULL violation; getting them half right is silence on every screen
    that reads a flat question row without joining its section."""
    parsed = _parsed(_section("A", "About", _question("How long?", 2)))
    _run(workbook.create_from_parsed(parsed))

    for call in _calls(world, "questionnairequestion", "create"):
        assert call["data"]["sectionCode"] == "A"
        assert call["data"]["sectionTitle"] == "About"
        assert call["data"]["questionnaireId"]
        assert call["data"]["sectionId"]


def test_no_answer_row_is_ever_written_on_the_create_path(world):
    parsed = _parsed(
        _section("A", "About", _question("How long?", 2, answers={"Ramesh": "Forty years"})),
        entryLabels=["Ramesh"],
    )
    _run(workbook.create_from_parsed(parsed))
    assert not [
        1 for model, op, _ in world.journal if model == "questionnaireresponse" and op != "find_many"
    ]


# =================================================================================================
# Re-uploading over an existing questionnaire
# =================================================================================================


def test_a_section_rename_updates_every_question_under_it(world):
    """R4, PINNED. ``sectionCode``/``sectionTitle`` are denormalised NOT NULL columns. Forgetting
    them on ``create`` is a loud NOT NULL violation; forgetting the ``update_many`` on rename is
    SILENT — last year's heading in the consolidated export, in the dataset CSV, on the browse rows,
    and in the ``?sectionCode=`` filter, with nothing wrong in the database to find."""
    parsed = _parsed(
        _section("A", "About the craft", _question("How long have you practised this craft?", 2)),
        _section("B", "Materials", _question("Where do you buy your yarn?", 3)),
    )
    _run(workbook.apply_parsed_edit("qnr-1", parsed))

    renames = _calls(world, "questionnairequestion", "update_many")
    assert len(renames) == 1, renames
    assert renames[0]["where"] == {"sectionId": "sec-a"}
    assert set(renames[0]["data"]) == {"sectionCode", "sectionTitle"}
    assert renames[0]["data"]["sectionTitle"] == "About the craft"


def test_every_existing_section_is_negated_before_any_positive_sort_order_is_written(world):
    """R1, PINNED AS CALL ORDER. ``@@unique([questionnaireId, sortOrder])`` means renumbering in
    place collides the instant a workbook swaps two sections — and it collides HALF WAY THROUGH,
    because there is no transaction anywhere in this backend. The negation pass has to finish before
    any positive number is claimed, including by a CREATE: a workbook whose first section is new asks
    for sortOrder 1 while an untouched section still holds it."""
    parsed = _parsed(
        _section("B", "Materials", _question("Where do you buy your yarn?", 2)),
        _section("A", "About", _question("How long have you practised this craft?", 3)),
    )
    _run(workbook.apply_parsed_edit("qnr-1", parsed))

    writes = [
        (index, kwargs["data"].get("sortOrder"))
        for index, (model, op, kwargs) in enumerate(world.journal)
        if model == "questionnairesection"
        and op in ("update", "create")
        and isinstance(kwargs.get("data", {}).get("sortOrder"), int)
    ]
    negatives = [index for index, value in writes if value < 0]
    positives = [index for index, value in writes if value > 0]
    assert negatives and positives
    assert max(negatives) < min(positives), writes


def test_no_section_is_left_at_a_negative_sort_order(world):
    """R2, AND IT SCANS EVERY SECTION RATHER THAN THE ONES THE WORKBOOK MOVED.

    The negation pass writes ``-position`` to the database. If the loop then guarded its write on
    "did the position change?" while reading the STALE pre-negation object, it would skip exactly the
    sections that did not move and strand them at a negative number for ever — sorting in front of
    everything on every read, and colliding with the NEXT upload's negation pass so an admin sees a
    failure caused by an upload made a month earlier. A test that checked only the moved sections
    would pass against that bug, which is why this one builds a workbook whose order is UNCHANGED and
    which drops a section entirely."""
    world.sections.append(
        SimpleNamespace(
            id="sec-c", questionnaireId="qnr-1", code="C", title="Selling", sortOrder=3, isActive=True
        )
    )
    world.sections.append(
        SimpleNamespace(
            id="sec-d", questionnaireId="qnr-1", code="D", title="Tools", sortOrder=4, isActive=True
        )
    )
    parsed = _parsed(
        _section("A", "About", _question("How long have you practised this craft?", 2)),
        _section("B", "Materials", _question("Where do you buy your yarn?", 3)),
        _section("C", "Selling"),
    )
    _run(workbook.apply_parsed_edit("qnr-1", parsed))

    final: dict[str, int] = {}
    for model, op, kwargs in world.journal:
        if model != "questionnairesection" or op != "update":
            continue
        if isinstance(kwargs["data"].get("sortOrder"), int):
            final[kwargs["where"]["id"]] = kwargs["data"]["sortOrder"]
    assert set(final) == {"sec-a", "sec-b", "sec-c", "sec-d"}, final
    assert all(value > 0 for value in final.values()), final


def test_a_section_absent_from_the_workbook_is_never_deleted(world):
    """§4.7(e), AND THE REASON IS INVISIBLE FROM THIS FILE.

    ``QuestionnaireSectionStatus.section`` is ON DELETE CASCADE and those rows are ADMIN COMPLETION
    VERDICTS — one explicit human judgement per (artisan, section) on the Check completion matrix.
    The ``answered`` set is built from ``QuestionnaireResponse`` and knows nothing about them, so a
    section can carry dozens of COMPLETED / NEEDS_REDO verdicts and zero text answers: the port's
    "safe to delete, nothing was ever answered" test says yes and every verdict is gone, permanently,
    with nothing on any screen saying so. This fake cannot model a Postgres cascade, so the only
    thing it CAN assert is the absence of the call — which is what it asserts, for a section with
    answered questions AND for one with none.

    Two more things hang off the same delete. ``QuestionnaireQuestion.sectionId`` is Restrict here,
    so the delete would in fact RAISE mid-loop; and ``AssignedTask.sectionIds`` has no foreign key at
    all, so live tasks would quietly lose a section while their denominator kept counting it."""
    _answered(world, "q-3")  # sec-b has an answered question; sec-a's are untouched
    world.sections.append(
        SimpleNamespace(
            id="sec-c", questionnaireId="qnr-1", code="C", title="Selling", sortOrder=3, isActive=True
        )
    )
    parsed = _parsed(_section("A", "About", _question("How long have you practised this craft?", 2)))
    _run(workbook.apply_parsed_edit("qnr-1", parsed))

    assert _calls(world, "questionnairesection", "delete") == []
    deactivations = {
        kwargs["where"]["id"]: kwargs["data"]
        for kwargs in _calls(world, "questionnairesection", "update")
        if kwargs["data"].get("isActive") is False
    }
    assert set(deactivations) == {"sec-b", "sec-c"}
    for data in deactivations.values():
        assert data["isActive"] is False
        assert data["sortOrder"] > 0


def test_a_retired_section_reports_what_still_points_at_it(world):
    """The admin is told what they have just switched off, in numbers they can act on. Neither count
    is enforced by anything: the verdicts survive because the delete does not happen, and the tasks
    are NOT rewritten — rewriting somebody's assigned work from inside a spreadsheet import is a
    decision this module has no standing to make."""
    world.questionnairesectionstatus.counts = 3
    world.assignedtask.counts = 2
    parsed = _parsed(_section("A", "About", _question("How long have you practised this craft?", 2)))
    report = _run(workbook.apply_parsed_edit("qnr-1", parsed))

    detail = next(d for d in report["details"] if d["action"] == "sectionRetired")
    assert "3 completion verdict" in detail["reason"]
    assert "2 assigned task" in detail["reason"]
    assert report["assignedTasksAffected"] == 2
    assert report["sectionsRetired"] == 1


def test_rewording_an_answered_question_creates_then_retires(world):
    """RULE 3. The recorded answers keep the wording they were given under; the new wording becomes a
    new question in the same place. CREATE FIRST, because the retirement has to name the
    replacement's id."""
    _answered(world, "q-1")
    parsed = _parsed(
        _section(
            "A",
            "About",
            _question("How many years have you worked this craft?", 2, questionId="q-1"),
            _question("Who taught you?", 3, questionId="q-2"),
        ),
        _section("B", "Materials", _question("Where do you buy your yarn?", 4, questionId="q-3")),
    )
    report = _run(workbook.apply_parsed_edit("qnr-1", parsed))

    creates = _calls(world, "questionnairequestion", "create")
    assert [c["data"]["prompt"] for c in creates] == ["How many years have you worked this craft?"]
    retire = next(
        c for c in _calls(world, "questionnairequestion", "update") if c["where"]["id"] == "q-1"
    )
    assert retire["data"]["isActive"] is False
    assert retire["data"]["retiredAt"] is not None
    assert retire["data"]["supersededById"]
    # ORDER: the create is journalled before the retire, or the retirement has no id to point at.
    ops = [(m, o, k) for m, o, k in world.journal if m == "questionnairequestion"]
    create_at = next(i for i, (_, o, k) in enumerate(ops) if o == "create")
    retire_at = next(
        i for i, (_, o, k) in enumerate(ops) if o == "update" and k["where"]["id"] == "q-1"
    )
    assert create_at < retire_at

    assert _calls(world, "questionnairequestion", "delete") == []
    assert report["superseded"] == 1
    assert report["versionAfter"] == report["versionBefore"] + 1


def test_rewording_an_unanswered_question_is_a_plain_update(world):
    """RULE 1. Nothing to protect, so nothing is duplicated: one update, no create, no retirement,
    and the version does NOT move — the number counts edits made after answers existed."""
    parsed = _parsed(
        _section(
            "A",
            "About",
            _question("How many years?", 2, questionId="q-1"),
            _question("Who taught you?", 3, questionId="q-2"),
        ),
        _section("B", "Materials", _question("Where do you buy your yarn?", 4, questionId="q-3")),
    )
    report = _run(workbook.apply_parsed_edit("qnr-1", parsed))

    assert _calls(world, "questionnairequestion", "create") == []
    updates = [
        c for c in _calls(world, "questionnairequestion", "update") if c["where"]["id"] == "q-1"
    ]
    assert len(updates) == 1
    assert updates[0]["data"]["prompt"] == "How many years?"
    assert "retiredAt" not in updates[0]["data"]
    assert report["superseded"] == 0
    assert report["versionAfter"] == report["versionBefore"]


def test_deleting_an_answered_question_retires_it(world):
    """RULE 4, and the server's own sentence is what the client prints. A question count cannot say
    "your answers are safe"; this can."""
    _answered(world, "q-2")
    parsed = _parsed(
        _section("A", "About", _question("How long have you practised this craft?", 2, questionId="q-1")),
        _section("B", "Materials", _question("Where do you buy your yarn?", 3, questionId="q-3")),
    )
    report = _run(workbook.apply_parsed_edit("qnr-1", parsed))

    retire = next(
        c for c in _calls(world, "questionnairequestion", "update") if c["where"]["id"] == "q-2"
    )
    assert retire["data"] == {"isActive": False, "retiredAt": retire["data"]["retiredAt"]}
    assert retire["data"]["retiredAt"] is not None
    assert _calls(world, "questionnairequestion", "delete") == []
    assert report["retired"] == 1
    detail = next(d for d in report["details"] if d["action"] == "retired")
    assert "retired rather than deleted" in detail["reason"]


def test_deleting_an_unanswered_question_really_deletes_it(world):
    """The other half of rule 1: a form littered with every question its author thought better of is
    not a form."""
    parsed = _parsed(
        _section("A", "About", _question("How long have you practised this craft?", 2, questionId="q-1")),
        _section("B", "Materials", _question("Where do you buy your yarn?", 3, questionId="q-3")),
    )
    report = _run(workbook.apply_parsed_edit("qnr-1", parsed))

    assert [c["where"]["id"] for c in _calls(world, "questionnairequestion", "delete")] == ["q-2"]
    assert report["removed"] == 1


def test_a_blank_answer_row_does_not_count_as_answered(world):
    """A RESPONSE SAVED AS A SINGLE SPACE IS NULL-ISH TO A PERSON AND NON-NULL TO POSTGRES, and only
    Python knows that — which is why the filter is in Python rather than in the where clause. The app
    writes exactly such a row when somebody opens an interview, tabs through it and saves, and
    treating that as "answered" would freeze a question nobody ever answered."""
    _answered(world, "q-1", text="   ")
    parsed = _parsed(
        _section(
            "A",
            "About",
            _question("How many years?", 2, questionId="q-1"),
            _question("Who taught you?", 3, questionId="q-2"),
        ),
        _section("B", "Materials", _question("Where do you buy your yarn?", 4, questionId="q-3")),
    )
    report = _run(workbook.apply_parsed_edit("qnr-1", parsed))

    assert report["superseded"] == 0
    updates = [
        c for c in _calls(world, "questionnairequestion", "update") if c["where"]["id"] == "q-1"
    ]
    assert updates[0]["data"]["prompt"] == "How many years?"


def test_a_note_with_no_answer_text_still_counts_as_answered(world):
    """A DIVERGENCE FROM THE PORT, AND IT IS DELIBERATE. This repository's ``QuestionnaireResponse``
    has a ``notes`` column the designer's answer table does not. An interviewer who wrote "artisan
    declined, see audio" in the notes and left the answer blank has recorded something ABOUT that
    wording, and rewording the question changes what their note appears to be a note about."""
    _answered(world, "q-1", text="", notes="Artisan declined; see the audio clip.")
    parsed = _parsed(
        _section(
            "A",
            "About",
            _question("How many years?", 2, questionId="q-1"),
            _question("Who taught you?", 3, questionId="q-2"),
        ),
        _section("B", "Materials", _question("Where do you buy your yarn?", 4, questionId="q-3")),
    )
    report = _run(workbook.apply_parsed_edit("qnr-1", parsed))
    assert report["superseded"] == 1


def test_a_question_with_only_blank_interview_rows_is_deactivated_rather_than_deleted(world):
    """A CONSTRAINT, NOT A RULE, AND THE PORT WOULD RAISE HERE.

    ``QuestionnaireResponse.question`` is ON DELETE RESTRICT, so ANY response row — including a
    completely blank one somebody left by opening an interview and saving — makes ``delete`` raise.
    The port has no second set and would issue that delete mid-loop, in a backend with no
    transaction: some questions deleted, some not, "upload failed" on screen, and no way for the
    admin to know what state their instrument is in. Deactivating is the honest answer, and deleting
    the interview rows to clear the way would be a write into somebody's fieldwork that a spreadsheet
    import has no business making."""
    _answered(world, "q-2", text="  ", notes="  ")
    parsed = _parsed(
        _section("A", "About", _question("How long have you practised this craft?", 2, questionId="q-1")),
        _section("B", "Materials", _question("Where do you buy your yarn?", 3, questionId="q-3")),
    )
    _run(workbook.apply_parsed_edit("qnr-1", parsed))

    assert _calls(world, "questionnairequestion", "delete") == []
    off = next(
        c for c in _calls(world, "questionnairequestion", "update") if c["where"]["id"] == "q-2"
    )
    assert off["data"] == {"isActive": False}


def test_a_question_switched_off_by_the_builder_is_reactivated_by_the_workbook(world):
    """§4.7(b2), FIRST HALF. ``isActive = false`` with ``retiredAt`` NULL means a person switched it
    off through the builder — DELETE /questionnaire/questions/{id} writes exactly that, and so does
    deleting a section. Naming it in the workbook is the admin saying it is part of the instrument
    again, so it comes back AND the row is applied normally afterwards."""
    world.questions[1].isActive = False
    world.questions[1].retiredAt = None
    parsed = _parsed(
        _section(
            "A",
            "About",
            _question("How long have you practised this craft?", 2, questionId="q-1"),
            _question("Who first taught you?", 3, questionId="q-2"),
        ),
        _section("B", "Materials", _question("Where do you buy your yarn?", 4, questionId="q-3")),
    )
    _run(workbook.apply_parsed_edit("qnr-1", parsed))

    update = next(
        c for c in _calls(world, "questionnairequestion", "update") if c["where"]["id"] == "q-2"
    )
    assert update["data"]["isActive"] is True
    assert update["data"]["prompt"] == "Who first taught you?"


def test_a_retired_question_is_not_reactivated_and_a_reworded_one_is_warned_about(world):
    """§4.7(b2), SECOND HALF. ``retiredAt`` set means the answer rule retired it. Reactivating on
    sight would mean that downloading a questionnaire and uploading it back UNCHANGED resurrects
    every question anybody ever replaced, each one standing next to its replacement, in the
    instrument forty researchers are answering that week."""
    world.questions[1].isActive = False
    world.questions[1].retiredAt = "2026-03-01T00:00:00Z"
    world.questions[1].supersededById = "q-9"
    parsed = _parsed(
        _section(
            "A",
            "About",
            _question("How long have you practised this craft?", 2, questionId="q-1"),
            _question("Who first taught you?", 9, questionId="q-2"),
        ),
        _section("B", "Materials", _question("Where do you buy your yarn?", 4, questionId="q-3")),
    )
    report = _run(workbook.apply_parsed_edit("qnr-1", parsed))

    touched = [
        c for c in _calls(world, "questionnairequestion", "update") if c["where"]["id"] == "q-2"
    ]
    assert touched == [], touched
    assert report["unchanged"] >= 1
    problem = next(p for p in report["problems"] if p["row"] == 9)
    assert problem["severity"] == "warning"
    assert "q-9" in problem["reason"]


def test_an_unknown_question_id_imports_as_new_and_reports_it(world):
    """An id from another instrument, or from a question that has since gone. Honouring it would
    graft another questionnaire's question onto this one; refusing the whole upload over one stale
    cell would be worse."""
    parsed = _parsed(
        _section(
            "A",
            "About",
            _question("How long have you practised this craft?", 2, questionId="q-1"),
            _question("Who taught you?", 3, questionId="q-2"),
            _question("A brand new question", 4, questionId="q-from-another-form"),
        ),
        _section("B", "Materials", _question("Where do you buy your yarn?", 5, questionId="q-3")),
    )
    report = _run(workbook.apply_parsed_edit("qnr-1", parsed))

    creates = _calls(world, "questionnairequestion", "create")
    assert [c["data"]["prompt"] for c in creates] == ["A brand new question"]
    problem = next(p for p in report["problems"] if "does not belong" in p["reason"])
    assert problem["row"] == 4
    assert problem["severity"] == "warning"


def test_the_prompt_fallback_matches_active_questions_only(world):
    """MATCHING STEP 2. A retired question keeps its original wording for ever, so a re-used wording
    would otherwise be silently matched onto the retired row — reactivating a question deliberately
    replaced, next to its replacement."""
    world.questions[1].isActive = False
    world.questions[1].retiredAt = "2026-03-01T00:00:00Z"
    parsed = _parsed(
        _section(
            "A",
            "About",
            _question("How long have you practised this craft?", 2),
            _question("Who taught you?", 3),  # the RETIRED question's exact wording, no id
        ),
        _section("B", "Materials", _question("Where do you buy your yarn?", 4)),
    )
    _run(workbook.apply_parsed_edit("qnr-1", parsed))

    creates = _calls(world, "questionnairequestion", "create")
    assert [c["data"]["prompt"] for c in creates] == ["Who taught you?"]
    assert not [
        c for c in _calls(world, "questionnairequestion", "update") if c["where"]["id"] == "q-2"
    ]


def test_an_orphaned_question_is_reported_and_never_touched(world):
    """§4.7(a1). ``sectionId`` is NULLABLE in this repository (the designer's is not), so the port's
    ``where={"sectionId": {"in": [...]}}`` read cannot see an orphan: never matched, never retired,
    never deleted, never reported — a live row, for ever, and an undeletable one if it has answers.
    The read here is by ``questionnaireId``, which is NOT NULL, so the orphan is seen and SAID.

    NOTHING CREATES ONE TODAY: ``QuestionnaireQuestion.section`` is Restrict since the
    questionnaire-instruments migration, so a section delete is refused rather than orphaning. This
    covers rows that predate that change, and the next change that reintroduces a way to make one.
    Deciding which section an orphan belongs to is a judgement about the contents of a research
    instrument, and this module does not make those — it reports and stops."""
    world.questions.append(
        SimpleNamespace(
            id="q-orphan",
            questionnaireId="qnr-1",
            sectionId=None,
            sectionCode="Z",
            sectionTitle="Gone",
            prompt="A question with no section",
            helpText=None,
            isRequired=False,
            sortOrder=1,
            isActive=True,
            retiredAt=None,
            supersededById=None,
        )
    )
    parsed = _parsed(
        _section(
            "A",
            "About",
            _question("How long have you practised this craft?", 2, questionId="q-1"),
            _question("Who taught you?", 3, questionId="q-2"),
        ),
        _section("B", "Materials", _question("Where do you buy your yarn?", 4, questionId="q-3")),
    )
    report = _run(workbook.apply_parsed_edit("qnr-1", parsed))

    problem = next(p for p in report["problems"] if "q-orphan" in p["reason"])
    assert problem["severity"] == "warning"
    assert not [
        c
        for c in _calls(world, "questionnairequestion", "update")
        + _calls(world, "questionnairequestion", "delete")
        if c.get("where", {}).get("id") == "q-orphan"
    ]


def test_answers_in_the_workbook_are_counted_reported_and_never_written(world):
    """R9, PINNED. The module this was ported from has a branch that IMPORTS the answers when the
    workbook looks hand-filled, and a reviewer diffing the two files will notice its absence and
    reinstate it. The result would be interview rows with invented artisan sets, attributed to the
    uploading admin, flowing into the consolidated export and the dataset API.

    The responses delegate is READ once — ``_answer_evidence`` needs it — and never written."""
    parsed = _parsed(
        _section(
            "A",
            "About",
            _question(
                "How long have you practised this craft?",
                2,
                questionId="q-1",
                answers={"Ramesh": "Forty years"},
                answerNotes={"Ramesh": "Said through his son"},
            ),
            _question("Who taught you?", 3, questionId="q-2", answers={"Ramesh": "My father"}),
        ),
        _section("B", "Materials", _question("Where do you buy your yarn?", 4, questionId="q-3")),
        entryLabels=["Ramesh"],
    )
    report = _run(workbook.apply_parsed_edit("qnr-1", parsed))

    writes = [
        op
        for model, op, _ in world.journal
        if model in ("questionnaireresponse", "questionnaireinterview") and op != "find_many"
    ]
    assert writes == [], writes
    assert report["answersSkipped"] == workbook._uploaded_answer_rows(parsed) == 2
    assert report["provenance"]["action"] == "answersNotImported"
    provenance = next(
        p for p in report["problems"] if p["reason"] == report["provenance"]["reason"]
    )
    assert provenance["severity"] == "warning"


REPORT_KEYS = {
    "created",
    "updated",
    "sections",
    "sectionsCreated",
    "sectionsRetired",
    "superseded",
    "retired",
    "removed",
    "unchanged",
    "answersSkipped",
    "assignedTasksAffected",
    "provenance",
    "versionBefore",
    "versionAfter",
    "problems",
    "details",
}


def test_the_report_carries_every_key_the_client_renders(world):
    """A DROPPED KEY IS A BLANK PANEL WITH NO ERROR, so this asserts the exact set against a literal
    and does it on BOTH paths. The client reads ``report.superseded`` and ``report.details``
    unconditionally; a report that omitted either would make "nothing was superseded"
    indistinguishable from "this client is reading a shape it does not understand"."""
    parsed = _parsed(_section("A", "About", _question("How long?", 2)))
    _, created = _run(workbook.create_from_parsed(parsed))
    edited = _run(workbook.apply_parsed_edit("qnr-1", parsed))
    assert set(created) == REPORT_KEYS
    assert set(edited) == REPORT_KEYS


def test_sections_means_the_same_thing_on_both_paths(world):
    """N9. The port used one key for "total parsed sections" on create and "newly created sections"
    on edit — one key, two meanings, one renderer, and an admin re-uploading an unchanged file
    reading "0 sections". ``sections`` is the workbook's count on both; ``sectionsCreated`` is the
    one that differs."""
    parsed = _parsed(
        _section("A", "About", _question("How long have you practised this craft?", 2)),
        _section("B", "Materials", _question("Where do you buy your yarn?", 3)),
    )
    _, created = _run(workbook.create_from_parsed(parsed))
    edited = _run(workbook.apply_parsed_edit("qnr-1", parsed))

    assert created["sections"] == edited["sections"] == 2
    assert created["sectionsCreated"] == 2
    assert edited["sectionsCreated"] == 0


# =================================================================================================
# The builder beside the workbook — the three quiet halves of one schema change
# =================================================================================================
#
# §6a adds `helpText` and `isRequired` to QuestionnaireQuestionCreate / …Update so the builder can
# set what the workbook sets. The schema change alone is only the LOUD half: without it, a POST
# carrying helpText is a 422 and somebody notices in five seconds. The three below are the quiet
# ones, and each is a different lie told with a 2xx.


def test_a_created_question_keeps_its_help_text_and_required_flag(world):
    """SILENT HALF ONE: a 201 that stores nothing.

    `create_question` builds its data dict FIELD BY FIELD while `update_question` beside it uses
    `model_dump(exclude_unset=True)`. Add the fields to the schema and stop there and the PATCH
    stores help text while the POST does not — two halves of one editor disagreeing, with no error
    anywhere and no way to tell from the response."""
    payload = workbook.QuestionnaireQuestionCreate(
        sectionId="sec-a",
        prompt="How is the yarn dyed?",
        helpText="Ask about the vat, not the colour.",
        isRequired=True,
    )
    _run(workbook.create_question(payload))

    data = _calls(world, "questionnairequestion", "create")[0]["data"]
    assert data["helpText"] == "Ask about the vat, not the colour."
    assert data["isRequired"] is True


def test_clearing_help_text_through_the_editor_actually_clears_it(world):
    """SILENT HALF TWO: a 200 that clears nothing.

    `clean_data` strips every None except the relation FKs in `CLEARABLE_KEYS`, and its own comment
    says that list is FK-only because blanking a SCALAR is governed per route. So `{"helpText": null}`
    — an admin deleting the guidance under a question — is stripped on the way in and the row keeps
    the old text behind a 200. A save that reports success and changes nothing is the worst of the
    three possible answers, because the person watching has no reason to look again."""
    world.questions[0].helpText = "Something that should go away."
    payload = workbook.QuestionnaireQuestionUpdate(helpText=None)
    _run(workbook.update_question("q-1", payload))

    data = _calls(world, "questionnairequestion", "update")[-1]["data"]
    assert "helpText" in data, data
    assert data["helpText"] is None


def test_omitting_help_text_from_a_patch_leaves_it_alone(world):
    """The other side of the line above, and the reason it is `model_fields_set` rather than a
    truthiness test: "the caller sent null" and "the caller said nothing" are different instructions,
    and conflating them would wipe the help text off every question a builder renames."""
    world.questions[0].helpText = "Keep me."
    payload = workbook.QuestionnaireQuestionUpdate(prompt="A reworded question")
    _run(workbook.update_question("q-1", payload))

    data = _calls(world, "questionnairequestion", "update")[-1]["data"]
    assert "helpText" not in data, data


# =================================================================================================
# The whole loop, through the real writers and the real parser
# =================================================================================================


def test_a_pro_forma_filled_in_uploaded_downloaded_and_uploaded_again_is_stable(world):
    """THE PRODUCT, END TO END, WITH ONLY THE DATABASE FAKED.

    Every other test in this file drives `create_from_parsed` / `apply_parsed_edit` from a
    hand-built `ParsedQuestionnaire`, which is the right shape for asserting one rule at a time and
    the wrong shape for catching a mismatch BETWEEN the writers and the parser — the two halves that
    have to be exact inverses of each other for the feature to work at all. A column heading the
    writer spells one way and the parser reads another would pass every test above and lose every
    question in production.

    So this runs the real loop: fill in a pro-forma the way an admin does, parse it, store it,
    export it back to a workbook through `export_instrument_payload`, parse THAT, and re-apply it.
    The property is that the second upload changes nothing — no create, no supersede, no retire —
    because an admin who downloads a questionnaire and uploads it back untouched must get their
    questionnaire back untouched.
    """
    from io import BytesIO

    from openpyxl import load_workbook

    from app.services.questionnaire_xlsx import (
        SHEET_QUESTIONS,
        build_pro_forma,
        build_questionnaire_workbook,
        parse_questionnaire_workbook,
    )

    # 1. An admin fills in the blank pro-forma, under its own headings, in Excel.
    book = load_workbook(BytesIO(build_pro_forma()))
    sheet = book[SHEET_QUESTIONS]
    typed = [
        ("A", "About the craft", "How long have you practised this craft?", "Count from the first unsupervised year.", "Yes"),
        ("", "", "Who taught you?", "", "No"),
        ("B", "Materials", "Where do you buy your yarn?", "", "Yes"),
    ]
    for offset, (code, title, prompt, help_text, required) in enumerate(typed, start=2):
        sheet.cell(row=offset, column=1, value=code)
        sheet.cell(row=offset, column=2, value=title)
        sheet.cell(row=offset, column=4, value=prompt)
        sheet.cell(row=offset, column=5, value=help_text)
        sheet.cell(row=offset, column=6, value=required)
    filled = BytesIO()
    book.save(filled)

    # 2. They upload it.
    parsed = parse_questionnaire_workbook(filled.getvalue(), filename="craft-toolkit.xlsx")
    new_id, created = _run(
        workbook.create_from_parsed(parsed, title="3rd Craft Toolkit Workshop", created_by_id="admin-1")
    )
    assert created["sections"] == 2
    assert created["created"] == 3
    assert [p for p in created["problems"] if p["severity"] == "error"] == []

    stored = [
        c["data"]
        for model, op, c in world.journal
        if model == "questionnairequestion" and op == "create"
    ]
    assert [row["prompt"] for row in stored] == [t[2] for t in typed]
    assert stored[0]["helpText"] == "Count from the first unsupervised year."
    assert [row["isRequired"] for row in stored] == [True, False, True]

    # 3. They download it again and upload it back UNCHANGED.
    payload = _run(workbook.export_instrument_payload(new_id))
    assert payload is not None
    assert payload["entry_labels"] == [], "the download must never carry an answer column per sitting"
    round_tripped = parse_questionnaire_workbook(build_questionnaire_workbook(**payload))
    assert round_tripped.questionnaireId == new_id

    world.journal.clear()
    report = _run(workbook.apply_parsed_edit(new_id, round_tripped))

    # 4. NOTHING MOVED. This is the assertion the whole test exists for.
    assert (report["created"], report["superseded"], report["retired"], report["removed"]) == (0, 0, 0, 0)
    assert report["sectionsRetired"] == 0
    assert report["unchanged"] == 3
    assert report["versionAfter"] == report["versionBefore"]
    assert _calls(world, "questionnairequestion", "create") == []
    assert _calls(world, "questionnairequestion", "delete") == []
    assert _calls(world, "questionnairesection", "delete") == []
