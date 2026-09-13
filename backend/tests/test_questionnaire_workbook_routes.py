"""The five workbook routes, driven over HTTP, with the database faked.

WHY THIS EXISTS BESIDE `test_questionnaire_workbook_import.py`. That file drives the import
FUNCTIONS directly, which is the right shape for asserting the edit rule and the wrong shape for
everything between the socket and them. None of the following is reachable from a direct call:

* whether `title` and `description` are FORM fields or query parameters — left as bare defaults
  FastAPI reads them off the query string and a client that posts them in the body gets an untitled
  questionnaire with a 201 saying it went fine;
* the suffix refusal, the size ceiling and the empty-upload refusal in `_read_workbook_upload`;
* whether `parse_questionnaire_workbook`'s sentence reaches the caller as a 422 `detail` or as a
  500 out of openpyxl;
* whether the 409 for a workbook downloaded from a DIFFERENT questionnaire fires before anything is
  written — which is the refusal standing between an admin picking the wrong file out of their
  downloads folder and every question in this instrument being switched off in one press.

THE DATABASE IS A STUB AND THE AUTH DEPENDENCY IS OVERRIDDEN, exactly as `test_permission_matrix.py`
does: the modules do `from app.core.db import db`, so each holds its own reference and rebinding the
module attribute by identity is what reaches them. Nothing here asserts a permission — that file
owns the gate — so every request below is made as an admin.
"""

import asyncio
import sys
from io import BytesIO
from types import SimpleNamespace
from typing import Any

import httpx
import pytest
from fastapi import FastAPI
from openpyxl import load_workbook

import app.core.db as core_db
from app.api.router import api_router
from app.core import deps
from app.services.questionnaire_xlsx import (
    _OLE2_MAGIC,
    SHEET_QUESTIONS,
    build_pro_forma,
)

ADMIN = SimpleNamespace(
    id="admin-1",
    email="admin@example.test",
    name="Admin",
    role="ADMIN",
    canManageQuestionnaire=False,
)

_APP = FastAPI()
_APP.include_router(api_router)
_APP.dependency_overrides[deps.get_current_user] = lambda: ADMIN


def _filled_pro_forma() -> bytes:
    """A pro-forma with three questions typed into it, the way an admin sends one."""
    book = load_workbook(BytesIO(build_pro_forma()))
    sheet = book[SHEET_QUESTIONS]
    rows = [
        ("A", "About the craft", "How long have you practised this craft?", "Yes"),
        ("", "", "Who taught you?", "No"),
        ("B", "Materials", "Where do you buy your yarn?", "maybe"),
    ]
    for offset, (code, title, prompt, required) in enumerate(rows, start=2):
        sheet.cell(row=offset, column=1, value=code)
        sheet.cell(row=offset, column=2, value=title)
        sheet.cell(row=offset, column=4, value=prompt)
        sheet.cell(row=offset, column=6, value=required)
    out = BytesIO()
    book.save(out)
    return out.getvalue()


class _Delegate:
    """Enough of a Prisma delegate for the happy path, recording nothing it is not asked for."""

    def __init__(self, rows: list[Any] | None = None) -> None:
        self.rows = rows or []
        self._next = 0

    async def create(self, data: dict[str, Any]) -> Any:
        self._next += 1
        row = SimpleNamespace(
            id=data.get("id") or f"made-{self._next}",
            version=1,
            isActive=True,
            isDefault=False,
            sortOrder=data.get("sortOrder", 1),
            sourceFilename=data.get("sourceFilename"),
            description=data.get("description"),
            createdAt="2026-09-13",
            **{k: v for k, v in data.items() if k not in ("id", "sortOrder", "sourceFilename", "description")},
        )
        self.rows.append(row)
        return row

    async def update(self, where: dict[str, Any], data: dict[str, Any]) -> Any:
        for row in self.rows:
            if row.id == where.get("id"):
                for key, value in data.items():
                    setattr(row, key, value)
                return row
        return SimpleNamespace(id=where.get("id"), **data)

    async def update_many(self, **_: Any) -> int:
        return 0

    async def delete(self, **_: Any) -> None:
        return None

    async def count(self, **_: Any) -> int:
        return 0

    async def find_unique(self, where: dict[str, Any], **_: Any) -> Any:
        return next((row for row in self.rows if row.id == where.get("id")), None)

    async def find_many(self, where: dict[str, Any] | None = None, **kwargs: Any) -> list[Any]:
        rows = list(self.rows)
        for key, clause in (where or {}).items():
            if isinstance(clause, dict) and "in" in clause:
                rows = [r for r in rows if getattr(r, key, None) in clause["in"]]
            else:
                rows = [r for r in rows if getattr(r, key, None) == clause]
        return rows[: kwargs["take"]] if kwargs.get("take") else rows


class _Db:
    def __init__(self, **delegates: _Delegate) -> None:
        self._delegates = delegates

    def __getattr__(self, name: str) -> Any:
        try:
            return self._delegates[name]
        except KeyError as exc:  # pragma: no cover - only read out of a failure
            raise AssertionError(f"the route reached for db.{name}, which this test did not model") from exc


@pytest.fixture
def stub_db(monkeypatch: pytest.MonkeyPatch) -> _Db:
    instrument = SimpleNamespace(
        id="qnr-1",
        title="2nd Craft Toolkit Workshop",
        description=None,
        version=4,
        sortOrder=1,
        isActive=True,
        isDefault=True,
        sourceFilename=None,
        createdAt="2026-01-01",
    )
    db = _Db(
        questionnaire=_Delegate([instrument]),
        questionnairesection=_Delegate(),
        questionnairequestion=_Delegate(),
        questionnaireresponse=_Delegate(),
        questionnairesectionstatus=_Delegate(),
        assignedtask=_Delegate(),
        workshop=_Delegate(),
    )
    real_db = core_db.db
    monkeypatch.setattr(core_db, "db", db)
    for module in list(sys.modules.values()):
        if getattr(module, "__name__", "").startswith("app.") and getattr(module, "db", None) is real_db:
            monkeypatch.setattr(module, "db", db)
    return db


def _post(path: str, *, files: Any = None, data: Any = None) -> httpx.Response:
    async def run() -> httpx.Response:
        transport = httpx.ASGITransport(app=_APP)
        async with httpx.AsyncClient(transport=transport, base_url="http://routes.test") as client:
            return await client.post(f"/api{path}", files=files, data=data)

    return asyncio.run(run())


def _get(path: str) -> httpx.Response:
    async def run() -> httpx.Response:
        transport = httpx.ASGITransport(app=_APP)
        async with httpx.AsyncClient(transport=transport, base_url="http://routes.test") as client:
            return await client.get(f"/api{path}")

    return asyncio.run(run())


XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


# --------------------------------------------------------------------------------------------
# Downloads
# --------------------------------------------------------------------------------------------


def test_the_pro_forma_downloads_as_a_workbook_with_a_filename_on_it(stub_db: _Db) -> None:
    """A download the browser cannot SAVE is a download that did not happen: without
    Content-Disposition the .xlsx is rendered as a page of mojibake in a new tab."""
    response = _get("/questionnaires/pro-forma")

    assert response.status_code == 200, response.text
    assert response.headers["content-type"] == XLSX
    assert 'filename="questionnaire-pro-forma.xlsx"' in response.headers["content-disposition"]
    assert response.content.startswith(b"PK\x03\x04")


def test_a_download_for_a_questionnaire_that_is_gone_is_a_404_not_a_500(stub_db: _Db) -> None:
    for path in ("/questionnaires/nope/xlsx", "/questionnaires/nope/question-set.xlsx"):
        response = _get(path)
        assert response.status_code == 404, (path, response.text)


# --------------------------------------------------------------------------------------------
# Uploads: the three refusals, cheapest first
# --------------------------------------------------------------------------------------------


def test_a_docx_is_refused_on_its_name_before_the_body_is_touched(stub_db: _Db) -> None:
    """THE CHEAPEST REFUSAL. The filename is a string on the multipart part, so a Word document is
    turned away without openpyxl being asked to open it — and the message names the remedy rather
    than saying "invalid file"."""
    response = _post(
        "/questionnaires/upload",
        files={"file": ("instrument.docx", b"PK\x03\x04not-a-workbook", "application/msword")},
    )

    assert response.status_code == 415, response.text
    detail = response.json()["detail"]
    assert "instrument.docx" in detail
    assert "Save As" in detail


def test_an_empty_upload_is_a_422_that_says_what_to_attach(stub_db: _Db) -> None:
    response = _post("/questionnaires/upload", files={"file": ("empty.xlsx", b"", XLSX)})

    assert response.status_code == 422, response.text
    assert "pro-forma" in response.json()["detail"]


def test_an_old_xls_renamed_to_xlsx_gets_the_save_as_message(stub_db: _Db) -> None:
    """THE SUFFIX CHECK CANNOT CATCH THIS ONE, which is the whole reason the parser looks at the
    file's own magic bytes. Somebody who renamed report.xls to report.xlsx to get past an upload
    filter has an .xls with an .xlsx name, and "invalid file" would send them nowhere."""
    response = _post(
        "/questionnaires/upload",
        files={"file": ("instrument.xlsx", _OLE2_MAGIC + b"\x00" * 4096, XLSX)},
    )

    assert response.status_code == 422, response.text
    detail = response.json()["detail"]
    assert "Save As" in detail
    assert "password" in detail


def test_a_workbook_with_no_question_column_is_a_422_naming_the_column(stub_db: _Db) -> None:
    from openpyxl import Workbook

    book = Workbook()
    book.active["A1"] = "Name"
    book.active["B1"] = "Notes"
    book.active["A2"] = "Ramesh"
    buffer = BytesIO()
    book.save(buffer)

    response = _post("/questionnaires/upload", files={"file": ("notes.xlsx", buffer.getvalue(), XLSX)})

    assert response.status_code == 422, response.text
    assert "Question" in response.json()["detail"]


# --------------------------------------------------------------------------------------------
# Uploads: the happy path, and the two things about it that are easy to get wrong
# --------------------------------------------------------------------------------------------


def test_an_uploaded_pro_forma_creates_a_questionnaire_and_reports_every_row_it_assumed_about(
    stub_db: _Db,
) -> None:
    """THE REPORT IS THE FEATURE. "maybe" in the Required column is the likeliest thing to be wrong
    with a real workbook and the one an admin cannot see from a question count, so the row comes back
    with its Excel row number and the question is imported as optional."""
    response = _post(
        "/questionnaires/upload",
        files={"file": ("craft-toolkit.xlsx", _filled_pro_forma(), XLSX)},
    )

    assert response.status_code == 201, response.text
    body = response.json()
    assert body["report"]["created"] == 3
    assert body["report"]["sections"] == 2
    assert body["questionnaire"]["id"]
    problem = next(p for p in body["report"]["problems"] if "yes/no" in p["reason"])
    assert problem["row"] == 4
    assert problem["severity"] == "warning"


def test_the_title_and_description_are_read_off_the_multipart_body_not_the_query_string(
    stub_db: _Db,
) -> None:
    """`Form(default=None)` IS WHAT MAKES THEM BODY FIELDS. Left as bare `str | None = None` defaults
    FastAPI reads them as QUERY parameters, and a client that posts them in the body has them
    silently ignored — an untitled questionnaire with a 201 saying it went fine. There is no pydantic
    model on a multipart route to carry the declaration instead, so this is the only guard."""
    response = _post(
        "/questionnaires/upload",
        files={"file": ("craft-toolkit.xlsx", _filled_pro_forma(), XLSX)},
        data={"title": "3rd Craft Toolkit Workshop", "description": "81 questions, 22 sections"},
    )

    assert response.status_code == 201, response.text
    created = stub_db.questionnaire.rows[-1]
    assert created.title == "3rd Craft Toolkit Workshop"
    assert created.description == "81 questions, 22 sections"
    # And the filename is kept, so two workbooks a fortnight apart are tellable apart on screen.
    assert created.sourceFilename == "craft-toolkit.xlsx"


def test_a_workbook_downloaded_from_another_questionnaire_is_a_409_before_anything_is_written(
    stub_db: _Db,
) -> None:
    """THE REFUSAL BETWEEN AN ADMIN AND THE WRONG FILE IN THEIR DOWNLOADS FOLDER.

    A Questionnaire ID on the Details sheet naming a DIFFERENT questionnaire means they picked the
    wrong one. Applying it would switch off this questionnaire's ENTIRE question set as "absent from
    the upload" in one press — every question, because none of the ids in the file belongs here. So
    the check runs before the first write, and the assertion is that NOTHING was created."""
    from app.services.questionnaire_xlsx import build_questionnaire_workbook

    other = build_questionnaire_workbook(
        title="A different instrument",
        description=None,
        questionnaire_id="qnr-SOMETHING-ELSE",
        version=2,
        sections=[
            {
                "code": "A",
                "title": "About",
                "questions": [
                    {"id": "q-x", "prompt": "How long?", "helpText": None, "isRequired": False,
                     "answers": {}, "answerNotes": {}}
                ],
            }
        ],
        entry_labels=[],
    )

    response = _post("/questionnaires/qnr-1/upload", files={"file": ("wrong.xlsx", other, XLSX)})

    assert response.status_code == 409, response.text
    assert "qnr-SOMETHING-ELSE" in response.json()["detail"]
    assert stub_db.questionnairesection.rows == []
    assert stub_db.questionnairequestion.rows == []


def test_a_reupload_over_a_questionnaire_that_is_gone_is_a_404(stub_db: _Db) -> None:
    response = _post(
        "/questionnaires/nope/upload", files={"file": ("craft-toolkit.xlsx", _filled_pro_forma(), XLSX)}
    )
    assert response.status_code == 404, response.text


def test_an_oversized_workbook_is_a_413_that_names_the_limit(stub_db: _Db) -> None:
    """THE CEILING, EXERCISED THROUGH THE REAL MULTIPART STACK RATHER THAN AGAINST THE HELPER.

    `tests/test_upload_bounds.py` proves `read_upload_bounded` refuses; what it cannot prove is that
    this route WIRES it up — passes `request` so the declared Content-Length is consulted, and passes
    `WORKBOOK_MAX_UPLOAD_BYTES` rather than some other number. A body just over the ceiling is sent
    for real, so starlette spools it and the framework's own parser runs, exactly as it would in
    production.

    Nine megabytes is deliberate: comfortably over 8 MiB plus `CONTENT_LENGTH_SLACK_BYTES`, and small
    enough that this test costs a second rather than a minute.
    """
    oversized = b"PK\x03\x04" + b"\x00" * (9 * 1024 * 1024)

    response = _post("/questionnaires/upload", files={"file": ("huge.xlsx", oversized, XLSX)})

    assert response.status_code == 413, response.status_code
    detail = response.json()["detail"]
    assert "8 MB" in detail
    assert "questionnaire workbook" in detail
    # Nothing was written: the refusal happens before the parser is ever handed the bytes.
    assert stub_db.questionnaire.rows[-1].id == "qnr-1"
